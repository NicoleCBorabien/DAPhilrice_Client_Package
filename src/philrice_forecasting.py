from pathlib import Path
from datetime import datetime
import re
import warnings

import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
from statsmodels.tsa.arima.model import ARIMA
from statsmodels.tsa.holtwinters import Holt

warnings.filterwarnings("ignore")

SPECIFIC_TAGS = [
    "core research function",
    "core administration function",
    "extension services",
    "ad hoc",
    "Research Enablement",
    "RSS",
    "SMART Farm",
    "RCEF",
    "RiceBIS",
    "Malusog Rice",
    "Special Project Support",
]

WINDOWS = {
    "5-Year": {"observations": 60, "validation": 12},
    "1-Year": {"observations": 12, "validation": 4},
    "6-Month": {"observations": 6, "validation": 2},
}

ARIMA_ORDERS = [
    (p, d, q)
    for p in range(3)
    for d in range(2)
    for q in range(3)
    if p + q <= 2
]

REQUIRED_COLUMNS = [
    "Employee",
    "Window",
    "Month",
    "Specific Tag",
    "% of Total Workload",
]

MONTH_SUMMARY_RE = re.compile(
    r"^(January|February|March|April|May|June|July|August|September|October|November|December)\s+(20\d{2})\s+Taxonomy Summary$"
)

WINDOW_ALIASES = {
    "5-year": "5-Year", "5 year": "5-Year", "5y": "5-Year", "60-month": "5-Year", "60 month": "5-Year",
    "1-year": "1-Year", "1 year": "1-Year", "1y": "1-Year", "12-month": "1-Year", "12 month": "1-Year",
    "6-month": "6-Month", "6 month": "6-Month", "6m": "6-Month",
}

TAG_LOOKUP = {t.casefold(): t for t in SPECIFIC_TAGS}


def rmse(actual, predicted):
    """Calculate root mean squared error between actual and predicted values."""
    actual = np.asarray(actual, dtype=float)
    predicted = np.asarray(predicted, dtype=float)
    return float(np.sqrt(np.mean((actual - predicted) ** 2)))


def normalize_window(value):
    """Normalize a user-supplied training-window label to the supported canonical window name."""
    if pd.isna(value):
        raise ValueError("Window is blank.")
    text = str(value).strip()
    if text in WINDOWS:
        return text
    normalized = WINDOW_ALIASES.get(text.casefold())
    if normalized is None:
        raise ValueError(
            f"Unsupported Window '{value}'. Use 5-Year, 1-Year, or 6-Month."
        )
    return normalized


def normalize_specific_tag(value):
    """Normalize a user-supplied Specific Tag name to the canonical tag used by the model."""
    if pd.isna(value):
        raise ValueError("Specific Tag is blank.")
    text = str(value).strip()
    tag = TAG_LOOKUP.get(text.casefold())
    if tag is None:
        raise ValueError(
            f"Unknown Specific Tag '{value}'. Allowed tags: "
            + ", ".join(SPECIFIC_TAGS)
        )
    return tag


def normalize_share(value):
    """Accept either Excel percentage fractions (0.083 = 8.3%) or plain 0-100 percentages (8.3 = 8.3%)."""
    if pd.isna(value):
        raise ValueError("% of Total Workload is blank.")
    x = float(value)
    if x < 0:
        raise ValueError("% of Total Workload cannot be negative.")
    if x <= 1.01:
        return x
    if x <= 100.1:
        return x / 100.0
    raise ValueError(f"Invalid % of Total Workload value: {value}")


def parse_month(value):
    # Pandas usually returns Excel dates as datetime values. Some workbook
    # writers expose the underlying Excel serial number instead, so support
    # both representations.
    """Convert a month value, including Excel serial dates, into a month-start pandas Timestamp."""
    if isinstance(value, (int, float, np.integer, np.floating)) and not pd.isna(value):
        if float(value) >= 1000:
            dt = pd.to_datetime(float(value), unit="D", origin="1899-12-30", errors="coerce")
        else:
            dt = pd.to_datetime(value, errors="coerce")
    else:
        dt = pd.to_datetime(value, errors="coerce")
    if pd.isna(dt):
        raise ValueError(f"Could not parse Month value '{value}'.")
    return pd.Timestamp(dt).to_period("M").to_timestamp()


def parse_standard_input_workbook(workbook_path):
    """
    Parse the preferred universal input format.

    Any sheet is accepted if it contains these columns:
      Employee | Window | Month | Specific Tag | % of Total Workload

    Returns an empty DataFrame if no sheet follows the standard format.
    """
    workbook_path = Path(workbook_path)
    records = []
    xls = pd.ExcelFile(workbook_path)

    ignored_sheets = {
        "Examples",
        "Examples - Replace Me",
        "Instructions",
        "Specific Tags",
    }

    for sheet in xls.sheet_names:
        if sheet.strip() in ignored_sheets:
            continue
        try:
            df = pd.read_excel(workbook_path, sheet_name=sheet, dtype=object)
        except Exception:
            continue

        # Normalize column whitespace while preserving display names.
        df.columns = [str(c).strip() for c in df.columns]
        if not all(col in df.columns for col in REQUIRED_COLUMNS):
            continue

        df = df[REQUIRED_COLUMNS].dropna(how="all")
        if df.empty:
            continue

        for row_num, row in df.iterrows():
            if all(pd.isna(row[col]) for col in REQUIRED_COLUMNS):
                continue
            try:
                employee = str(row["Employee"]).strip()
                if not employee or employee.lower() == "nan":
                    raise ValueError("Employee is blank.")
                records.append({
                    "Employee": employee,
                    "Window": normalize_window(row["Window"]),
                    "Month": parse_month(row["Month"]),
                    "Specific Tag": normalize_specific_tag(row["Specific Tag"]),
                    "% Total Workload": normalize_share(row["% of Total Workload"]),
                    "Source File": workbook_path.name,
                    "Source Sheet": sheet,
                })
            except Exception as exc:
                raise ValueError(
                    f"{workbook_path.name} / {sheet} / Excel row {row_num + 2}: {exc}"
                ) from exc

    return pd.DataFrame(records)


def discover_standard_inputs(source_dir):
    """Scan all .xlsx files in Source Data. If any contain standard-format rows, use those rows."""
    source_dir = Path(source_dir)
    frames = []
    for path in sorted(source_dir.glob("*.xlsx")):
        if path.name.startswith("~$"):
            continue
        frame = parse_standard_input_workbook(path)
        if not frame.empty:
            frames.append(frame)
    if not frames:
        return pd.DataFrame()
    return pd.concat(frames, ignore_index=True)


def build_matrix_from_records(records, employee, window):
    """Build and validate the monthly Specific Tag workload-share matrix for one employee and window."""
    subset = records[
        (records["Employee"] == employee)
        & (records["Window"] == window)
    ].copy()
    if subset.empty:
        return None

    # Reject duplicate tag rows for the same employee/window/month.
    dup = subset.duplicated(["Month", "Specific Tag"], keep=False)
    if dup.any():
        sample = subset.loc[dup, ["Month", "Specific Tag", "Source File"]].head(10)
        raise ValueError(
            f"Duplicate rows found for {employee} / {window}. "
            "Each Month + Specific Tag must appear once. Examples:\n"
            + sample.to_string(index=False)
        )

    pivot = subset.pivot(index="Month", columns="Specific Tag", values="% Total Workload").sort_index()
    for tag in SPECIFIC_TAGS:
        if tag not in pivot.columns:
            pivot[tag] = 0.0
    pivot = pivot[SPECIFIC_TAGS].fillna(0.0).astype(float)

    required_n = WINDOWS[window]["observations"]
    if len(pivot) < required_n:
        raise ValueError(
            f"{employee} / {window} has {len(pivot)} unique months. "
            f"At least {required_n} are required."
        )
    if len(pivot) > required_n:
        pivot = pivot.iloc[-required_n:].copy()

    # Require continuous monthly observations.
    expected = pd.date_range(pivot.index.min(), pivot.index.max(), freq="MS")
    if len(expected) != len(pivot) or not np.array_equal(expected.values, pivot.index.values):
        raise ValueError(
            f"{employee} / {window} must contain consecutive monthly observations with no missing months."
        )

    totals = pivot.sum(axis=1)
    if not np.allclose(totals.values, 1.0, atol=0.002):
        bad = totals[~np.isclose(totals.values, 1.0, atol=0.002)]
        details = ", ".join(f"{d.strftime('%b %Y')}={v*100:.2f}%" for d, v in bad.items())
        raise ValueError(
            f"{employee} / {window} monthly Specific Tag shares must total about 100%. Problem months: {details}"
        )

    # Correct only small source rounding differences.
    pivot = pivot.div(totals, axis=0)
    return pivot


def source_from_standard_records(records):
    """Create employee/window modeling matrices from standardized input records and derive shorter windows when needed."""
    employees = sorted(records["Employee"].dropna().unique().tolist())
    source = {}
    skipped = []
    for employee in employees:
        source[employee] = {}

        # First load every window explicitly supplied by the user.
        for window in WINDOWS:
            if not ((records["Employee"] == employee) & (records["Window"] == window)).any():
                continue
            try:
                source[employee][window] = build_matrix_from_records(records, employee, window)
            except ValueError as exc:
                skipped.append(str(exc))

        # Convenience derivation: explicit shorter windows override derived ones.
        # A user can therefore provide only a 5-Year series and automatically
        # receive 5-Year, 1-Year, and 6-Month views.
        if "5-Year" in source[employee] and "1-Year" not in source[employee]:
            source[employee]["1-Year"] = source[employee]["5-Year"].iloc[-12:].copy()
        if "6-Month" not in source[employee]:
            if "1-Year" in source[employee]:
                source[employee]["6-Month"] = source[employee]["1-Year"].iloc[-6:].copy()
            elif "5-Year" in source[employee]:
                source[employee]["6-Month"] = source[employee]["5-Year"].iloc[-6:].copy()

        if not source[employee]:
            source.pop(employee)

    if not source:
        raise ValueError("No valid employee/window datasets were found in the standard input files.")
    return source, skipped


# --------------------------
# Legacy PhilRice sample support
# --------------------------

def extract_employee_name_from_sheet(workbook_path, sheet_name):
    """Extract the employee display name from a legacy PhilRice worksheet heading."""
    df = pd.read_excel(workbook_path, sheet_name=sheet_name, header=None, dtype=object)
    pattern = re.compile(
        r"^(January|February|March|April|May|June|July|August|September|October|November|December)\s+20\d{2}\s+[—–-]\s+(.+)$"
    )
    for value in df.to_numpy().ravel():
        if isinstance(value, str):
            m = pattern.match(value.strip())
            if m:
                return m.group(2).strip()
    # Safe fallback to sheet name if a heading is unavailable.
    return sheet_name.replace("_", " ").strip()


def parse_legacy_monthly_summary(workbook_path, employee_sheet):
    """Parse monthly Specific Tag workload shares from a legacy 5-Year or 1-Year PhilRice worksheet."""
    df = pd.read_excel(workbook_path, sheet_name=employee_sheet, header=None, dtype=object)
    current_month = None
    records = []
    for _, row in df.iterrows():
        values = row.tolist()
        for value in values:
            if isinstance(value, str):
                match = MONTH_SUMMARY_RE.match(value.strip())
                if match:
                    current_month = datetime.strptime(
                        f"{match.group(1)} {match.group(2)}", "%B %Y"
                    )
                    break
        if current_month is None:
            continue
        for idx in range(len(values) - 2):
            general_tag, specific_tag, pct = values[idx:idx+3]
            if (
                general_tag in ("Functional", "Strategic Program")
                and isinstance(specific_tag, str)
                and specific_tag.strip().casefold() in TAG_LOOKUP
                and pd.notna(pct)
            ):
                records.append({
                    "Month": pd.Timestamp(current_month),
                    "Specific Tag": TAG_LOOKUP[specific_tag.strip().casefold()],
                    "% Total Workload": float(pct),
                })
                break
    return pd.DataFrame(records)


def parse_legacy_six_month(workbook_path, employee_sheet):
    """Parse July-to-December 2025 Specific Tag workload shares from a legacy taxonomy-mapping worksheet."""
    df = pd.read_excel(workbook_path, sheet_name=employee_sheet, header=None, dtype=object)
    current_month = None
    summary_col = None
    records = []
    month_title_re = re.compile(
        r"^(July|August|September|October|November|December)\s+2025\s+Taxonomy Summary$"
    )
    for _, row in df.iterrows():
        values = row.tolist()
        for idx, value in enumerate(values):
            if isinstance(value, str):
                match = month_title_re.match(value.strip())
                if match:
                    current_month = datetime.strptime(f"{match.group(1)} 2025", "%B %Y")
                    summary_col = idx
                    break
        if current_month is None or summary_col is None:
            continue

        # General Tag | Specific Tag | %
        if summary_col + 2 < len(values):
            g, tag, pct = values[summary_col:summary_col+3]
            if (
                g in ("Functional", "Strategic Program")
                and isinstance(tag, str)
                and tag.strip().casefold() in TAG_LOOKUP
                and pd.notna(pct)
            ):
                records.append({
                    "Month": pd.Timestamp(current_month),
                    "Specific Tag": TAG_LOOKUP[tag.strip().casefold()],
                    "% Total Workload": float(pct),
                })
                continue

        # Specific Tag | % (Lanuza-style summary)
        if summary_col + 2 < len(values):
            tag, pct = values[summary_col + 1:summary_col + 3]
            if (
                isinstance(tag, str)
                and tag.strip().casefold() in TAG_LOOKUP
                and pd.notna(pct)
            ):
                records.append({
                    "Month": pd.Timestamp(current_month),
                    "Specific Tag": TAG_LOOKUP[tag.strip().casefold()],
                    "% Total Workload": float(pct),
                })
    return pd.DataFrame(records)


def _legacy_matrix(df, required_n, employee, window):
    """Convert parsed legacy records into a validated monthly Specific Tag workload-share matrix."""
    if df.empty:
        raise ValueError(f"No legacy source rows found for {employee} / {window}.")
    pivot = df.pivot_table(index="Month", columns="Specific Tag", values="% Total Workload", aggfunc="sum", fill_value=0).sort_index()
    for tag in SPECIFIC_TAGS:
        if tag not in pivot.columns:
            pivot[tag] = 0.0
    pivot = pivot[SPECIFIC_TAGS].astype(float)
    if len(pivot) != required_n:
        raise ValueError(f"{employee} / {window}: expected {required_n} months, found {len(pivot)}.")
    totals = pivot.sum(axis=1)
    if not np.allclose(totals.values, 1.0, atol=0.002):
        raise ValueError(f"{employee} / {window}: source shares do not total approximately 100%.")
    return pivot.div(totals, axis=0)


def load_legacy_sample_sources(source_dir):
    """Load employee modeling matrices from the bundled legacy 5-Year, 1-Year, and 6-Month PhilRice source workbooks."""
    source_dir = Path(source_dir)
    five = source_dir / "PhilRice_5Year_Dummy.xlsx"
    one = source_dir / "PhilRice_1Year_Dummy.xlsx"
    six = source_dir / "Taxonomy_Mapping_6Month.xlsx"
    if not all(p.exists() for p in (five, one, six)):
        raise FileNotFoundError(
            "No filled standard-format input was found, and the legacy sample source files are incomplete. "
            "See README.md for the required universal input format."
        )

    five_sheets = [s for s in pd.ExcelFile(five).sheet_names if s not in ("READ ME", "Monthly Summary")]
    one_sheets = set(pd.ExcelFile(one).sheet_names)
    six_sheets = set(pd.ExcelFile(six).sheet_names)
    sheets = [s for s in five_sheets if s in one_sheets and s in six_sheets]
    if not sheets:
        raise ValueError("No matching employee sheets exist across the three legacy sample workbooks.")

    source = {}
    skipped = []
    for sheet in sheets:
        employee = extract_employee_name_from_sheet(five, sheet)
        try:
            source[employee] = {
                "5-Year": _legacy_matrix(parse_legacy_monthly_summary(five, sheet), 60, employee, "5-Year"),
                "1-Year": _legacy_matrix(parse_legacy_monthly_summary(one, sheet), 12, employee, "1-Year"),
                "6-Month": _legacy_matrix(parse_legacy_six_month(six, sheet), 6, employee, "6-Month"),
            }
        except Exception as exc:
            skipped.append(f"Skipped legacy sample employee {employee}: {exc}")
    if not source:
        raise ValueError("No valid employees were found across the three legacy sample workbooks.")
    return source, skipped


def load_source_folder(source_dir):
    """
    Main source entrypoint.

    Priority:
    1. Any filled workbook(s) following the universal five-column format.
    2. If none are present, the bundled three legacy PhilRice sample workbooks.
    """
    records = discover_standard_inputs(source_dir)
    if not records.empty:
        source, skipped = source_from_standard_records(records)
        return source, skipped, "standard"
    source, skipped = load_legacy_sample_sources(source_dir)
    return source, skipped, "legacy-sample"


# --------------------------
# Modeling
# --------------------------

def rolling_arima(series, order, validation_size):
    """Perform rolling one-step-ahead validation for a specified ARIMA order and return forecasts and RMSE."""
    y = np.asarray(series, dtype=float)
    history = list(y[:-validation_size])
    actual, forecast = [], []
    for observed in y[-validation_size:]:
        fit = ARIMA(history, order=order).fit()
        predicted = float(fit.forecast(1)[0])
        actual.append(float(observed))
        forecast.append(predicted)
        history.append(float(observed))
    return actual, forecast, rmse(actual, forecast)


def rolling_holt(series, validation_size):
    """Perform rolling one-step-ahead validation with Holt's linear trend model and return forecasts and RMSE."""
    y = np.asarray(series, dtype=float)
    history = list(y[:-validation_size])
    actual, forecast = [], []
    for observed in y[-validation_size:]:
        fit = Holt(history, initialization_method="estimated", damped_trend=False).fit(optimized=True)
        predicted = float(fit.forecast(1)[0])
        actual.append(float(observed))
        forecast.append(predicted)
        history.append(float(observed))
    return actual, forecast, rmse(actual, forecast)


def select_model(series, validation_size):
    """Select the best ARIMA candidate by validation RMSE and compare it with Holt's Trend."""
    y = np.asarray(series, dtype=float)
    if np.allclose(y, 0):
        return {
            "best_arima_order": None,
            "arima_rmse": 0.0,
            "holt_rmse": 0.0,
            "selected_model": "Constant zero",
            "selected_rmse": 0.0,
            "validation_actual": list(map(float, y[-validation_size:])),
            "validation_forecast": [0.0] * validation_size,
        }

    best = None
    for order in ARIMA_ORDERS:
        try:
            actual, forecast, error = rolling_arima(y, order, validation_size)
            if best is None or error < best["rmse"]:
                best = {"order": order, "actual": actual, "forecast": forecast, "rmse": error}
        except Exception:
            continue

    try:
        h_actual, h_forecast, h_error = rolling_holt(y, validation_size)
    except Exception:
        h_actual = list(map(float, y[-validation_size:]))
        h_forecast = [float(y[-validation_size - 1])] * validation_size
        h_error = rmse(h_actual, h_forecast)

    if best is None or h_error < best["rmse"]:
        return {
            "best_arima_order": None if best is None else best["order"],
            "arima_rmse": float("inf") if best is None else best["rmse"],
            "holt_rmse": h_error,
            "selected_model": "Holt's Trend",
            "selected_rmse": h_error,
            "validation_actual": h_actual,
            "validation_forecast": h_forecast,
        }

    return {
        "best_arima_order": best["order"],
        "arima_rmse": best["rmse"],
        "holt_rmse": h_error,
        "selected_model": f"ARIMA{best['order']}",
        "selected_rmse": best["rmse"],
        "validation_actual": best["actual"],
        "validation_forecast": best["forecast"],
    }


def refit_and_forecast(series, selection, horizon=12):
    """Refit the selected model on the full training series and generate the requested forecast horizon."""
    y = np.asarray(series, dtype=float)
    if selection["selected_model"] == "Constant zero":
        return np.zeros(horizon, dtype=float)
    if selection["selected_model"].startswith("ARIMA"):
        fit = ARIMA(y, order=selection["best_arima_order"]).fit()
        future = np.asarray(fit.forecast(horizon), dtype=float)
    else:
        fit = Holt(y, initialization_method="estimated", damped_trend=False).fit(optimized=True)
        future = np.asarray(fit.forecast(horizon), dtype=float)
    return np.maximum(future, 0.0)


def model_one_window(employee, window, matrix, horizon=12):
    """Validate, select, refit, and forecast all Specific Tags for one employee and training window."""
    validation_size = WINDOWS[window]["validation"]
    tag_results, raw_future = [], []
    validation_dates = [d.strftime("%B %Y") for d in matrix.index[-validation_size:]]

    for tag in SPECIFIC_TAGS:
        series = matrix[tag].to_numpy(dtype=float)
        selection = select_model(series, validation_size)
        future = refit_and_forecast(series, selection, horizon)
        raw_future.append(future)
        tag_results.append({
            "employee": employee,
            "window": window,
            "specific_tag": tag,
            "observations": len(series),
            "validation_months": validation_size,
            "best_arima_order": selection["best_arima_order"],
            "arima_rmse": selection["arima_rmse"],
            "holt_rmse": selection["holt_rmse"],
            "selected_model": selection["selected_model"],
            "selected_rmse": selection["selected_rmse"],
            "validation_dates": validation_dates,
            "validation_actual": selection["validation_actual"],
            "validation_forecast": selection["validation_forecast"],
        })

    raw = np.vstack(raw_future).T
    normalized = raw.copy()
    for i in range(normalized.shape[0]):
        total = float(normalized[i].sum())
        if total > 0:
            normalized[i] /= total
    for j, result in enumerate(tag_results):
        result["forecast"] = normalized[:, j].tolist()

    forecast_dates = pd.date_range(
        matrix.index[-1] + pd.offsets.MonthBegin(1),
        periods=horizon,
        freq="MS",
    )

    return {
        "employee": employee,
        "window": window,
        "tag_results": tag_results,
        "forecast_dates": [d.strftime("%B %Y") for d in forecast_dates],
        "forecast_month_totals": normalized.sum(axis=1).tolist(),
    }


def run_forecasting_from_folder(source_dir, horizon=12):
    """Load source data from a folder and run forecasting for every valid employee and available window."""
    source, skipped, input_mode = load_source_folder(source_dir)
    results = {}
    for employee, windows in source.items():
        results[employee] = {}
        for window, matrix in windows.items():
            results[employee][window] = model_one_window(employee, window, matrix, horizon)
    return source, results, skipped, input_mode


def results_to_dataframes(results):
    """Convert nested modeling results into tabular DataFrames for summary, validation, forecast, and total-check outputs."""
    model_rows, validation_rows, forecast_rows, total_rows = [], [], [], []

    for employee in sorted(results):
        for window in WINDOWS:
            if window not in results[employee]:
                continue
            block = results[employee][window]
            for forecast_date, total in zip(block["forecast_dates"], block["forecast_month_totals"]):
                total_rows.append({
                    "Employee": employee,
                    "Window": window,
                    "Forecast Month": forecast_date,
                    "Total %": total,
                })
            for r in block["tag_results"]:
                model_rows.append({
                    "Employee": employee,
                    "Window": window,
                    "Specific Tag": r["specific_tag"],
                    "Obs.": r["observations"],
                    "Validation": r["validation_months"],
                    "Best ARIMA": "N/A" if r["best_arima_order"] is None else str(tuple(r["best_arima_order"])),
                    "ARIMA RMSE (pp)": r["arima_rmse"] * 100,
                    "Holt RMSE (pp)": r["holt_rmse"] * 100,
                    "Selected Model": r["selected_model"],
                    "Selected RMSE (pp)": r["selected_rmse"] * 100,
                })
                for date, actual, forecast in zip(r["validation_dates"], r["validation_actual"], r["validation_forecast"]):
                    validation_rows.append({
                        "Employee": employee,
                        "Window": window,
                        "Specific Tag": r["specific_tag"],
                        "Validation Month": date,
                        "Actual %": actual,
                        "Forecast %": forecast,
                        "Error (pp)": (forecast - actual) * 100,
                    })
                for forecast_date, value in zip(block["forecast_dates"], r["forecast"]):
                    forecast_rows.append({
                        "Employee": employee,
                        "Window": window,
                        "Specific Tag": r["specific_tag"],
                        "Forecast Month": forecast_date,
                        "Forecast %": value,
                    })

    return {
        "Model Summary": pd.DataFrame(model_rows),
        "Validation": pd.DataFrame(validation_rows),
        "12-Month Forecasts": pd.DataFrame(forecast_rows),
        "Forecast Total Check": pd.DataFrame(total_rows),
    }


def safe_sheet_name(employee, used):
    """Create a unique Excel-safe chart-sheet name for an employee."""
    base = re.sub(r"[\\/*?:\[\]]", "_", employee).strip() or "Employee"
    base = base[:24]
    candidate = f"{base} Charts"[:31]
    n = 2
    while candidate in used:
        suffix = f" {n}"
        candidate = (f"{base[:31-len(' Charts')-len(suffix)]}{suffix} Charts")[:31]
        n += 1
    used.add(candidate)
    return candidate



def write_results_excel(results, output_path, skipped=None, input_mode=None):
    """
    Create the client-facing Excel workbook.

    Charts are rendered with Matplotlib and embedded as PNG images rather than
    relying on native Excel chart-axis rendering. This guarantees that month
    tick labels and percentage tick labels remain visible in Excel.
    """
    import tempfile

    from matplotlib.ticker import PercentFormatter
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    wb.remove(wb.active)

    navy = "1F4E78"
    white = "FFFFFF"
    yellow = "FFF2CC"
    light = "D9EAF7"

    # --------------------------
    # README
    # --------------------------
    readme = wb.create_sheet("README")
    readme["A1"] = "PhilRice Specific Tag Forecasting Results"
    readme["A1"].font = Font(bold=True, size=18, color=white)
    readme["A1"].fill = PatternFill("solid", fgColor=navy)
    readme.merge_cells("A1:H1")

    readme["A3"] = (
        "Input mode: " + str(input_mode or "unknown") + ". "
        "The model reads source workbooks from 02_Source_Data and does not require "
        "hardcoded employee names or precomputed model results."
    )
    readme.merge_cells("A3:H5")
    readme["A3"].alignment = Alignment(wrap_text=True, vertical="top")

    readme["A7"] = (
        "Process: Prepare tag % series → Train ARIMA & Holt → Rolling validation → "
        "Calculate RMSE → Select model → Refit selected model → Forecast 12 months → "
        "Normalize monthly total to 100%."
    )
    readme.merge_cells("A7:H9")
    readme["A7"].alignment = Alignment(wrap_text=True, vertical="top")

    if skipped:
        readme["A11"] = "Input warnings / skipped datasets"
        readme["A11"].font = Font(bold=True)
        for i, msg in enumerate(skipped, start=12):
            readme.cell(i, 1, msg)

    # --------------------------
    # Result tables
    # --------------------------
    tables = results_to_dataframes(results)

    for sheet_name in [
        "Model Summary",
        "Validation",
        "12-Month Forecasts",
        "Forecast Total Check",
    ]:
        df = tables[sheet_name]
        ws = wb.create_sheet(sheet_name)
        ws.append(list(df.columns))

        for c in ws[1]:
            c.fill = PatternFill("solid", fgColor=navy)
            c.font = Font(bold=True, color=white)
            c.alignment = Alignment(horizontal="center", wrap_text=True)

        for row in df.itertuples(index=False, name=None):
            ws.append(list(row))

        ws.freeze_panes = "A2"

        for col_idx in range(1, ws.max_column + 1):
            ws.column_dimensions[get_column_letter(col_idx)].width = 16

        if ws.max_column >= 3:
            ws.column_dimensions["C"].width = 28

        if sheet_name == "Validation":
            for row in range(2, ws.max_row + 1):
                ws.cell(row, 5).number_format = "0.00%"
                ws.cell(row, 6).number_format = "0.00%"
                ws.cell(row, 7).number_format = "0.000"

        elif sheet_name == "12-Month Forecasts":
            for row in range(2, ws.max_row + 1):
                ws.cell(row, 5).number_format = "0.00%"

        elif sheet_name == "Forecast Total Check":
            for row in range(2, ws.max_row + 1):
                ws.cell(row, 4).number_format = "0.00%"

    used_sheet_names = set(wb.sheetnames)

    # --------------------------
    # Chart sheets
    # --------------------------
    with tempfile.TemporaryDirectory() as tmpdir:
        tmpdir = Path(tmpdir)

        for employee in sorted(results):
            chart_sheet_name = safe_sheet_name(employee, used_sheet_names)
            ws = wb.create_sheet(chart_sheet_name)

            ws["A1"] = f"{employee} - Forecast and Validation Charts"
            ws["A1"].font = Font(bold=True, size=16, color=white)
            ws["A1"].fill = PatternFill("solid", fgColor=navy)
            ws.merge_cells("A1:M1")

            ws["A2"] = (
                "All 11 Specific Tags are shown, including constant-zero tags. "
                "Charts are embedded images so axis labels display consistently."
            )
            ws.merge_cells("A2:M2")
            ws["A2"].fill = PatternFill("solid", fgColor=yellow)
            ws["A2"].alignment = Alignment(wrap_text=True)

            # Give the chart area predictable dimensions.
            for col in range(1, 14):
                ws.column_dimensions[get_column_letter(col)].width = 12

            # --------------------------
            # Forecast charts
            # --------------------------
            forecast_chart_idx = 0

            for window in WINDOWS:
                if window not in results[employee]:
                    continue

                block = results[employee][window]
                x_labels = block["forecast_dates"]

                fig, ax = plt.subplots(figsize=(12, 5.6))

                for r in block["tag_results"]:
                    ax.plot(
                        range(len(x_labels)),
                        np.asarray(r["forecast"], dtype=float),
                        marker="o",
                        linewidth=1.5,
                        markersize=3,
                        label=r["specific_tag"],
                    )

                ax.set_title(f"{window}: 12-Month Forecast", fontsize=12, pad=12)
                ax.set_xlabel("Forecast Month", fontsize=10)
                ax.set_ylabel("% of Total Workload", fontsize=10)

                ax.set_xticks(range(len(x_labels)))
                ax.set_xticklabels(
                    x_labels,
                    rotation=45,
                    ha="right",
                    fontsize=8,
                )

                ax.set_ylim(0, 1)
                ax.set_yticks(np.arange(0, 1.01, 0.10))
                ax.yaxis.set_major_formatter(PercentFormatter(xmax=1.0, decimals=0))
                ax.tick_params(axis="y", labelsize=8)

                ax.grid(True, axis="y", alpha=0.3)
                ax.legend(
                    loc="center left",
                    bbox_to_anchor=(1.01, 0.5),
                    fontsize=7,
                    frameon=False,
                )

                fig.subplots_adjust(
                    left=0.10,
                    right=0.73,
                    bottom=0.28,
                    top=0.88,
                )

                png = tmpdir / (
                    f"{re.sub(r'[^A-Za-z0-9]+', '_', employee)}_"
                    f"{window.replace('-', '')}_forecast.png"
                )
                fig.savefig(png, dpi=160, bbox_inches="tight")
                plt.close(fig)

                img = XLImage(str(png))
                img.width = 900
                img.height = 420

                anchor_row = 4 + forecast_chart_idx * 22
                ws.add_image(img, f"A{anchor_row}")
                forecast_chart_idx += 1

            # --------------------------
            # Validation section
            # --------------------------
            section_row = max(72, 4 + forecast_chart_idx * 22 + 1)

            ws[f"A{section_row}"] = "Validation Charts: Historical Actual vs Forecast"
            ws[f"A{section_row}"].font = Font(
                bold=True,
                size=14,
                color=white,
            )
            ws[f"A{section_row}"].fill = PatternFill("solid", fgColor=navy)
            ws.merge_cells(
                start_row=section_row,
                start_column=1,
                end_row=section_row,
                end_column=13,
            )

            validation_chart_idx = 0

            for window in WINDOWS:
                if window not in results[employee]:
                    continue

                by_tag = {
                    r["specific_tag"]: r
                    for r in results[employee][window]["tag_results"]
                }

                for tag in SPECIFIC_TAGS:
                    r = by_tag[tag]

                    actual = np.asarray(r["validation_actual"], dtype=float)
                    predicted = np.asarray(r["validation_forecast"], dtype=float)
                    x_labels = r["validation_dates"]
                    x = np.arange(len(x_labels))

                    fig, ax = plt.subplots(figsize=(6.4, 4.2))

                    ax.plot(
                        x,
                        actual,
                        marker="o",
                        linewidth=1.8,
                        markersize=4,
                        label="Actual",
                    )
                    ax.plot(
                        x,
                        predicted,
                        marker="o",
                        linewidth=1.8,
                        markersize=4,
                        label="Forecast",
                    )

                    ax.set_title(
                        f"{window}: {tag}\nHistorical Actual vs Forecast",
                        fontsize=9,
                        pad=8,
                    )
                    ax.set_xlabel("Validation Month", fontsize=9)
                    ax.set_ylabel("% of Total Workload", fontsize=9)

                    ax.set_xticks(x)
                    ax.set_xticklabels(
                        x_labels,
                        rotation=45,
                        ha="right",
                        fontsize=7,
                    )

                    all_values = np.concatenate([actual, predicted])
                    finite = all_values[np.isfinite(all_values)]

                    if finite.size:
                        low = min(float(finite.min()), 0.0)
                        high = max(float(finite.max()), 0.0)

                        if np.isclose(low, high):
                            low = min(-0.05, low - 0.05)
                            high = max(0.05, high + 0.05)
                        else:
                            pad = max((high - low) * 0.15, 0.02)
                            low -= pad
                            high += pad

                        ax.set_ylim(low, high)

                    ax.yaxis.set_major_formatter(
                        PercentFormatter(xmax=1.0, decimals=0)
                    )
                    ax.tick_params(axis="y", labelsize=7)
                    ax.grid(True, axis="y", alpha=0.3)
                    ax.legend(fontsize=7, frameon=False)

                    fig.subplots_adjust(
                        left=0.16,
                        right=0.97,
                        bottom=0.30,
                        top=0.82,
                    )

                    png = tmpdir / (
                        f"{re.sub(r'[^A-Za-z0-9]+', '_', employee)}_"
                        f"{window.replace('-', '')}_"
                        f"{re.sub(r'[^A-Za-z0-9]+', '_', tag)}_validation.png"
                    )
                    fig.savefig(png, dpi=160, bbox_inches="tight")
                    plt.close(fig)

                    img = XLImage(str(png))
                    img.width = 450
                    img.height = 315

                    anchor_row = section_row + 3 + (
                        validation_chart_idx // 2
                    ) * 18
                    anchor_col = "A" if validation_chart_idx % 2 == 0 else "G"

                    ws.add_image(img, f"{anchor_col}{anchor_row}")
                    validation_chart_idx += 1

            # Keep compact validation source tables on the right for auditability.
            support_col = 15
            support_row = 1

            for window in WINDOWS:
                if window not in results[employee]:
                    continue

                block = results[employee][window]
                ws.cell(support_row, support_col, f"{window} Forecast Month")

                for j, tag in enumerate(SPECIFIC_TAGS, start=1):
                    ws.cell(support_row, support_col + j, tag)

                by_tag = {
                    r["specific_tag"]: r
                    for r in block["tag_results"]
                }

                for i, month in enumerate(block["forecast_dates"], start=1):
                    ws.cell(support_row + i, support_col, month)

                    for j, tag in enumerate(SPECIFIC_TAGS, start=1):
                        cell = ws.cell(
                            support_row + i,
                            support_col + j,
                            by_tag[tag]["forecast"][i - 1],
                        )
                        cell.number_format = "0.0%"

                support_row += 16

            support_row = max(support_row, 52)

            for window in WINDOWS:
                if window not in results[employee]:
                    continue

                by_tag = {
                    r["specific_tag"]: r
                    for r in results[employee][window]["tag_results"]
                }

                for tag in SPECIFIC_TAGS:
                    r = by_tag[tag]
                    title = (
                        f"Validation: {employee} | {window} | {tag} | "
                        "Historical Actual vs Forecast"
                    )

                    ws.merge_cells(
                        start_row=support_row,
                        start_column=15,
                        end_row=support_row,
                        end_column=17,
                    )

                    title_cell = ws.cell(support_row, 15, title)
                    title_cell.fill = PatternFill("solid", fgColor=navy)
                    title_cell.font = Font(bold=True, color=white)
                    title_cell.alignment = Alignment(wrap_text=True)

                    header_row = support_row + 1

                    for j, value in enumerate(
                        ["Month", "Actual", "Forecast"],
                        start=15,
                    ):
                        cell = ws.cell(header_row, j, value)
                        cell.fill = PatternFill("solid", fgColor=light)
                        cell.font = Font(bold=True)
                        cell.alignment = Alignment(horizontal="center")

                    for i, (date, actual_value, forecast_value) in enumerate(
                        zip(
                            r["validation_dates"],
                            r["validation_actual"],
                            r["validation_forecast"],
                        ),
                        start=1,
                    ):
                        ws.cell(header_row + i, 15, date)
                        ws.cell(
                            header_row + i,
                            16,
                            actual_value,
                        ).number_format = "0.0%"
                        ws.cell(
                            header_row + i,
                            17,
                            forecast_value,
                        ).number_format = "0.0%"

                    support_row = header_row + len(
                        r["validation_dates"]
                    ) + 3

            for col in range(15, 27):
                ws.column_dimensions[get_column_letter(col)].width = 14

        wb.save(output_path)

    return output_path

def verify_results(results):
    """Validate core result invariants, including tag count, forecast horizon, nonnegative forecasts, and 100% monthly totals."""
    for employee, windows in results.items():
        for window, block in windows.items():
            assert len(block["tag_results"]) == len(SPECIFIC_TAGS)
            assert all(len(r["forecast"]) == 12 for r in block["tag_results"])
            assert all(min(r["forecast"]) >= 0 for r in block["tag_results"])
            assert all(abs(total - 1.0) < 1e-8 for total in block["forecast_month_totals"])
    return True


def default_package_paths(script_file=None):
    """Resolve the package root, source-data folder, and default results-workbook path from the script location."""
    here = Path(script_file or __file__).resolve().parent
    # Works when file is in 03_Model_and_Code or in package root.
    if here.name == "03_Model_and_Code":
        root = here.parent
    else:
        root = here
    return root, root / "02_Source_Data", root / "01_Results" / "PhilRice_Specific_Tag_Forecast_Results.xlsx"


if __name__ == "__main__":
    root, source_dir, output_file = default_package_paths(__file__)
    print(f"Source folder: {source_dir}")
    source, results, skipped, mode = run_forecasting_from_folder(source_dir)
    verify_results(results)
    write_results_excel(results, output_file, skipped=skipped, input_mode=mode)
    print(f"Input mode: {mode}")
    print(f"Employees modeled: {', '.join(sorted(results))}")
    for employee in sorted(results):
        print(f" - {employee}: {', '.join(results[employee].keys())}")
    if skipped:
        print("Warnings:")
        for msg in skipped:
            print(" -", msg)
    print(f"Results written to: {output_file}")
